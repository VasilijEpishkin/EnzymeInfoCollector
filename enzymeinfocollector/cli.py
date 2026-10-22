"""User-facing CLI for collecting and exporting enzyme information."""

from __future__ import annotations

import argparse
import csv
import json
import logging
import os
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Any

from .collector import EnzymeCollector
from .sources import DataSourceError, HttpCache, default_cache_dir, validate_ec_number

LOGGER = logging.getLogger("enzymeinfocollector")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="enzymeinfocollector",
        description="Collect curated enzyme nomenclature, UniProt proteins, and Rhea reactions.",
    )
    inputs = parser.add_mutually_exclusive_group(required=True)
    inputs.add_argument("--enzyme", action="append", help="Enzyme name; can be repeated")
    inputs.add_argument("--ec", action="append", help="Full EC number; can be repeated")
    inputs.add_argument("--file", type=Path, help="CSV, TSV, TXT, or XLSX input file")
    parser.add_argument("--output", "-o", type=Path, default=Path("enzyme_results.json"), help="Output file (default: enzyme_results.json)")
    parser.add_argument("--format", choices=("auto", "json", "csv", "fasta"), default="auto", help="Output format; auto uses the output extension")
    parser.add_argument("--cache-dir", type=Path, default=default_cache_dir(), help="HTTP cache directory")
    parser.add_argument("--contact-email", default=os.environ.get("ENZYMEINFOCOLLECTOR_EMAIL"), help="Contact email added to upstream User-Agent headers")
    parser.add_argument("--timeout", type=float, default=30.0, help="HTTP timeout in seconds (default: 30)")
    parser.add_argument("--retries", type=int, default=3, help="Retries for transient upstream errors (default: 3)")
    parser.add_argument("--refresh", action="store_true", help="Bypass cached responses")
    parser.add_argument("--verbose", action="store_true", help="Log progress and cache use to stderr")
    return parser


def read_queries(args: argparse.Namespace) -> list[tuple[str, str]]:
    if args.enzyme:
        queries = [("enzyme", value.strip()) for value in args.enzyme]
    elif args.ec:
        queries = [("ec", validate_ec_number(value)) for value in args.ec]
    else:
        queries = _read_query_file(args.file)
    if not queries or any(not value for _, value in queries):
        raise ValueError("Input must contain at least one non-empty enzyme name or EC number")
    unique: list[tuple[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for query in queries:
        key = (query[0], query[1].casefold())
        if key not in seen:
            unique.append(query)
            seen.add(key)
    return unique


def _read_query_file(path: Path | None) -> list[tuple[str, str]]:
    if path is None or not path.is_file():
        raise ValueError(f"Input file does not exist: {path}")
    suffix = path.suffix.casefold()
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise ValueError("XLSX input requires openpyxl; install it with `pip install enzymeinfocollector[xlsx]`") from error
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except (OSError, ValueError, zipfile.BadZipFile) as error:
            raise ValueError(f"Could not read Excel workbook {path}: {error}") from error
        try:
            rows = workbook.active.iter_rows(values_only=True)
            return _queries_from_rows(list(rows))
        finally:
            workbook.close()
    if suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            return _queries_from_rows(list(csv.reader(stream, delimiter=delimiter)))
    if suffix == ".txt":
        with path.open("r", encoding="utf-8-sig") as stream:
            values = [line.strip() for line in stream if line.strip() and not line.lstrip().startswith("#")]
            return [("ec", validate_ec_number(value)) if value.casefold().startswith("ec ") or (value[:1].isdigit() and value.count(".") >= 2) else ("enzyme", value) for value in values]
    raise ValueError("Input file must have .csv, .tsv, .txt, or .xlsx extension")


def _queries_from_rows(rows: list[tuple[Any, ...]] | list[list[Any]]) -> list[tuple[str, str]]:
    if not rows:
        return []
    header = [str(value or "").strip().casefold() for value in rows[0]]
    enzyme_columns = [index for index, name in enumerate(header) if name in {"protein", "enzyme", "enzyme name", "protein name"}]
    ec_columns = [index for index, name in enumerate(header) if name in {"ec", "ec number", "ec_number"}]
    if not enzyme_columns and not ec_columns:
        raise ValueError("Input table needs a Protein/Enzyme or EC/EC number column")
    queries = []
    for row in rows[1:]:
        for index in enzyme_columns:
            value = _cell(row, index)
            if value:
                queries.append(("enzyme", value))
        for index in ec_columns:
            value = _cell(row, index)
            if value:
                queries.append(("ec", validate_ec_number(value)))
    return queries


def _cell(row: tuple[Any, ...] | list[Any], index: int) -> str:
    return str(row[index]).strip() if index < len(row) and row[index] is not None else ""


def write_output(path: Path, output_format: str, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    selected_format = output_format
    if selected_format == "auto":
        suffix = path.suffix.casefold()
        selected_format = "csv" if suffix == ".csv" else "fasta" if suffix in {".fa", ".faa", ".fasta"} else "json"
    descriptor, temp_name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    os.close(descriptor)
    temp_path = Path(temp_name)
    try:
        if selected_format == "json":
            with temp_path.open("w", encoding="utf-8", newline="\n") as stream:
                json.dump(payload, stream, ensure_ascii=False, indent=2)
                stream.write("\n")
                stream.flush()
                os.fsync(stream.fileno())
        elif selected_format == "csv":
            _write_csv(temp_path, payload)
        else:
            _write_fasta(temp_path, payload)
        os.replace(temp_path, path)
    finally:
        temp_path.unlink(missing_ok=True)


CSV_FIELDS = (
    "input_type", "input_value", "status", "ec_number", "enzyme_name", "alternative_names",
    "catalytic_activity", "uniprot_accession", "entry_name", "protein_name", "gene_names",
    "organism", "organism_id", "reviewed", "length", "sequence", "rhea_id", "equation",
    "chebi_names", "chebi_ids", "reaction_smiles", "reaction_smiles_directions", "enzyme_source", "uniprot_source", "rhea_source", "rhea_smiles_source", "rhea_directions_source",
)


def _write_csv(path: Path, payload: dict[str, Any]) -> None:
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=CSV_FIELDS, extrasaction="ignore")
        writer.writeheader()
        for result in payload["results"]:
            if not result["enzymes"]:
                writer.writerow({
                    "input_type": result["input"]["type"],
                    "input_value": result["input"]["value"],
                    "status": result["status"],
                })
                continue
            for enzyme in result["enzymes"]:
                proteins = enzyme["proteins"] or [None]
                reactions = enzyme["reactions"] or [None]
                for protein in proteins:
                    for reaction in reactions:
                        writer.writerow({
                            "input_type": result["input"]["type"],
                            "input_value": result["input"]["value"],
                            "status": result["status"],
                            "ec_number": enzyme["ec_number"],
                            "enzyme_name": enzyme["enzyme_name"],
                            "alternative_names": _json_cell(enzyme["alternative_names"]),
                            "catalytic_activity": _json_cell(enzyme["catalytic_activity"]),
                            "uniprot_accession": protein["accession"] if protein else "",
                            "entry_name": protein["entry_name"] if protein else "",
                            "protein_name": protein["protein_name"] if protein else "",
                            "gene_names": _json_cell(protein["gene_names"]) if protein else "",
                            "organism": protein["organism"] if protein else "",
                            "organism_id": protein["organism_id"] if protein else "",
                            "reviewed": protein["reviewed"] if protein else "",
                            "length": protein["length"] if protein else "",
                            "sequence": protein["sequence"] if protein else "",
                            "rhea_id": f"RHEA:{reaction['rhea_id']}" if reaction else "",
                            "equation": reaction["equation"] if reaction else "",
                            "chebi_names": _json_cell(reaction["chebi_names"]) if reaction else "",
                            "chebi_ids": _json_cell(reaction["chebi_ids"]) if reaction else "",
                            "reaction_smiles": reaction["reaction_smiles"] if reaction and reaction["reaction_smiles"] else "",
                            "reaction_smiles_directions": _json_cell(reaction["reaction_smiles_directions"]) if reaction else "",
                            "enzyme_source": _json_cell(enzyme["provenance"]["enzyme_nomenclature"]),
                            "uniprot_source": _json_cell(enzyme["provenance"]["protein_annotations"]),
                            "rhea_source": _json_cell(enzyme["provenance"]["reactions"]),
                            "rhea_smiles_source": _json_cell(enzyme["provenance"]["reaction_smiles"]),
                            "rhea_directions_source": _json_cell(enzyme["provenance"]["reaction_directions"]),
                        })
        stream.flush()
        os.fsync(stream.fileno())


def _json_cell(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def _write_fasta(path: Path, payload: dict[str, Any]) -> None:
    written: set[str] = set()
    with path.open("w", encoding="utf-8", newline="\n") as stream:
        for result in payload["results"]:
            for enzyme in result["enzymes"]:
                for protein in enzyme["proteins"]:
                    accession = protein["accession"]
                    sequence = protein["sequence"]
                    if not sequence or accession in written:
                        continue
                    written.add(accession)
                    ec_numbers = ",".join(protein["ec_numbers"] or [enzyme["ec_number"]])
                    organism = "_".join(protein["organism"].split()) or "unknown"
                    stream.write(f">{accession} EC={ec_numbers} OS={organism} source=UniProtKB\n")
                    for offset in range(0, len(sequence), 80):
                        stream.write(sequence[offset:offset + 80] + "\n")
        stream.flush()
        os.fsync(stream.fileno())


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    logging.basicConfig(level=logging.INFO if args.verbose else logging.WARNING, format="%(levelname)s: %(message)s")
    if args.timeout <= 0 or args.retries < 0:
        parser.error("--timeout must be positive and --retries must be non-negative")
    try:
        queries = read_queries(args)
        if args.file and args.output.resolve() == args.file.resolve():
            raise ValueError("Output path must not overwrite the input file")
        cache = HttpCache(args.cache_dir, contact_email=args.contact_email, timeout=args.timeout, retries=args.retries)
        payload = EnzymeCollector(cache, refresh=args.refresh).collect(queries)
        write_output(args.output, args.format, payload)
        not_found = sum(result["status"] == "not_found" for result in payload["results"])
        LOGGER.info("Wrote %d input result(s) to %s (%d unmatched)", len(payload["results"]), args.output, not_found)
        if not args.verbose:
            print(f"Wrote {len(payload['results'])} result(s) to {args.output}; unmatched: {not_found}", file=sys.stderr)
        return 0
    except (ValueError, DataSourceError, OSError) as error:
        print(f"enzymeinfocollector: {error}", file=sys.stderr)
        return 1
